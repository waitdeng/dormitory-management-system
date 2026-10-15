"""Employee service — CRUD, detail, Excel two-phase import/export (M2)."""

from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import (
    ConflictError,
    ResourceNotFound,
    ValidationError,
)
from app.models.base import utcnow
from app.models.department import Department
from app.models.employee import Employee
from app.models.import_job import ImportJob
from app.schemas.employee import (
    EmployeeCreate,
    EmployeeUpdate,
)
from app.services import audit_service

# Spreadsheet formula-injection guard (spec §11)
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

from typing import TYPE_CHECKING  # noqa: E402

if TYPE_CHECKING:
    from app.services.data_scope import DeptScope


def _safe_excel_value(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _to_out(
    emp: Employee,
    dept: Department | None = None,
    *,
    accommodation_status: str | None = None,
) -> dict:
    return {
        "id": emp.id,
        "employee_no": emp.employee_no,
        "full_name": emp.full_name,
        "gender": emp.gender,
        "department_id": emp.department_id,
        "department_code": dept.department_code if dept else None,
        "department_name": dept.department_name if dept else None,
        "position": emp.position,
        "employment_type": emp.employment_type,
        "phone": emp.phone,
        "join_date": emp.join_date,
        "termination_date": emp.termination_date,
        "employment_status": emp.employment_status,
        "accommodation_required": emp.accommodation_required,
        "shift_type": emp.shift_type,
        "notes": emp.notes,
        # Derived from accommodation_records / bed_reservations (M3).
        "accommodation_status": accommodation_status or "not_in_dorm",
        "created_at": emp.created_at,
        "updated_at": emp.updated_at,
    }


async def _load(db: AsyncSession, employee_id: int) -> Employee:
    emp = (
        await db.execute(
            select(Employee)
            .where(Employee.id == employee_id)
            .execution_options(populate_existing=True)
        )
    ).scalar_one_or_none()
    if emp is None:
        raise ResourceNotFound(f"Employee {employee_id} not found")
    return emp


async def _dept_map(db: AsyncSession, dept_ids: set[int]) -> dict[int, Department]:
    if not dept_ids:
        return {}
    rows = (await db.execute(select(Department).where(Department.id.in_(dept_ids)))).scalars()
    return {d.id: d for d in rows}


async def _accommodation_status_map(
    db: AsyncSession, employee_ids: list[int]
) -> dict[int, str]:
    """Derive in_dorm / reserved / not_in_dorm from M3 relation tables."""
    if not employee_ids:
        return {}
    from app.models.accommodation import AccommodationRecord, BedReservation

    ids = list({int(i) for i in employee_ids})
    occupied = set(
        (
            await db.execute(
                select(AccommodationRecord.employee_id).where(
                    AccommodationRecord.employee_id.in_(ids),
                    AccommodationRecord.status == "active",
                )
            )
        ).scalars().all()
    )
    reserved = set(
        (
            await db.execute(
                select(BedReservation.employee_id).where(
                    BedReservation.employee_id.in_(ids),
                    BedReservation.status == "active",
                )
            )
        ).scalars().all()
    )
    out: dict[int, str] = {}
    for eid in ids:
        if eid in occupied:
            out[eid] = "in_dorm"
        elif eid in reserved:
            out[eid] = "reserved"
        else:
            out[eid] = "not_in_dorm"
    return out


async def list_employees(
    db: AsyncSession,
    *,
    page: int = 1,
    page_size: int = 20,
    keyword: str | None = None,
    department_id: int | None = None,
    employment_status: str | None = None,
    accommodation_status: str | None = None,
    accommodation_required: bool | None = None,
    scope: DeptScope | None = None,
) -> tuple[list[dict], int]:
    from app.services.data_scope import apply_employee_scope

    query = select(Employee)
    count_query = select(func.count()).select_from(Employee)

    if scope is not None:
        query = apply_employee_scope(query, scope, Employee.id)
        count_query = apply_employee_scope(count_query, scope, Employee.id)

    if keyword:
        pattern = f"%{keyword}%"
        cond = (Employee.employee_no.ilike(pattern)) | (Employee.full_name.ilike(pattern))
        query = query.where(cond)
        count_query = count_query.where(cond)
    if department_id is not None:
        query = query.where(Employee.department_id == department_id)
        count_query = count_query.where(Employee.department_id == department_id)
    if employment_status:
        query = query.where(Employee.employment_status == employment_status)
        count_query = count_query.where(Employee.employment_status == employment_status)
    if accommodation_required is not None:
        query = query.where(Employee.accommodation_required == accommodation_required)
        count_query = count_query.where(
            Employee.accommodation_required == accommodation_required
        )

    # BUG-014: visual-allocation candidate set — strict eligibility
    if accommodation_status == "allocatable":
        from app.models.accommodation import AccommodationRecord, BedReservation

        in_dorm = select(AccommodationRecord.employee_id).where(
            AccommodationRecord.status == "active"
        )
        reserved = select(BedReservation.employee_id).where(
            BedReservation.status == "active"
        )
        query = query.where(
            Employee.employment_status == "active",
            Employee.accommodation_required.is_(True),
            ~Employee.id.in_(in_dorm),
            ~Employee.id.in_(reserved),
        )
        count_query = count_query.where(
            Employee.employment_status == "active",
            Employee.accommodation_required.is_(True),
            ~Employee.id.in_(in_dorm),
            ~Employee.id.in_(reserved),
        )
    elif accommodation_status:
        from app.models.accommodation import AccommodationRecord, BedReservation

        if accommodation_status == "in_dorm":
            sub = select(AccommodationRecord.employee_id).where(
                AccommodationRecord.status == "active"
            )
            query = query.where(Employee.id.in_(sub))
            count_query = count_query.where(Employee.id.in_(sub))
        elif accommodation_status == "reserved":
            in_dorm = select(AccommodationRecord.employee_id).where(
                AccommodationRecord.status == "active"
            )
            reserved = select(BedReservation.employee_id).where(
                BedReservation.status == "active"
            )
            query = query.where(Employee.id.in_(reserved), ~Employee.id.in_(in_dorm))
            count_query = count_query.where(
                Employee.id.in_(reserved), ~Employee.id.in_(in_dorm)
            )
        elif accommodation_status == "not_in_dorm":
            in_dorm = select(AccommodationRecord.employee_id).where(
                AccommodationRecord.status == "active"
            )
            reserved = select(BedReservation.employee_id).where(
                BedReservation.status == "active"
            )
            query = query.where(
                ~Employee.id.in_(in_dorm), ~Employee.id.in_(reserved)
            )
            count_query = count_query.where(
                ~Employee.id.in_(in_dorm), ~Employee.id.in_(reserved)
            )
        else:
            return [], 0

    total = (await db.execute(count_query)).scalar_one()
    stmt = (
        query.order_by(Employee.id.desc()).offset((page - 1) * page_size).limit(page_size)
    )
    rows = (await db.execute(stmt)).scalars().all()
    depts = await _dept_map(db, {r.department_id for r in rows if r.department_id})
    acc_map = await _accommodation_status_map(db, [r.id for r in rows])
    items = [
        _to_out(
            r,
            depts.get(r.department_id) if r.department_id else None,
            accommodation_status=acc_map.get(r.id),
        )
        for r in rows
    ]
    return items, total


async def get_employee(
    db: AsyncSession, employee_id: int, *, scope: DeptScope | None = None
) -> dict:
    emp = await _load(db, employee_id)
    if scope is not None:
        scope.require(emp.department_id, resource=f"employee {employee_id}")
    depts = await _dept_map(db, {emp.department_id} if emp.department_id else set())
    acc_map = await _accommodation_status_map(db, [emp.id])
    return _to_out(
        emp,
        depts.get(emp.department_id) if emp.department_id else None,
        accommodation_status=acc_map.get(emp.id),
    )


async def create_employee(
    db: AsyncSession,
    payload: EmployeeCreate,
    *,
    actor_user_id: int | None = None,
    request_id: str | None = None,
) -> dict:
    existing = (
        await db.execute(select(Employee).where(Employee.employee_no == payload.employee_no))
    ).scalar_one_or_none()
    if existing is not None:
        raise ConflictError(f"Employee no '{payload.employee_no}' already exists")

    if payload.department_id is not None:
        await _assert_department_active(db, payload.department_id)
    _assert_dates(payload.join_date, payload.termination_date)

    emp = Employee(**payload.model_dump())
    db.add(emp)
    await db.flush()

    await audit_service.write_audit(
        db,
        action="employee.create",
        resource_type="employee",
        resource_id=emp.id,
        actor_user_id=actor_user_id,
        request_id=request_id,
        after_data=_to_out(emp),
    )
    await db.commit()
    return await get_employee(db, emp.id)


async def update_employee(
    db: AsyncSession,
    employee_id: int,
    payload: EmployeeUpdate,
    *,
    actor_user_id: int | None = None,
    request_id: str | None = None,
) -> dict:
    emp = await _load(db, employee_id)
    depts_before = await _dept_map(db, {emp.department_id} if emp.department_id else set())
    before = _to_out(emp, depts_before.get(emp.department_id) if emp.department_id else None)

    if payload.employee_no is not None and payload.employee_no != emp.employee_no:
        existing = (
            await db.execute(
                select(Employee).where(
                    Employee.employee_no == payload.employee_no,
                    Employee.id != employee_id,
                )
            )
        ).scalar_one_or_none()
        if existing is not None:
            raise ConflictError(f"Employee no '{payload.employee_no}' already exists")
        emp.employee_no = payload.employee_no

    if "department_id" in payload.model_fields_set:
        if payload.department_id is not None:
            await _assert_department_active(db, payload.department_id)
        emp.department_id = payload.department_id

    data = payload.model_dump(exclude_unset=True)
    # employment_status transitions to offboarding/terminated go through the
    # M4 event pipeline so todos/alerts are created (P0 A/B). Other fields
    # are applied first in this transaction.
    routed_status = data.pop("employment_status", None)
    for field in (
        "full_name",
        "gender",
        "position",
        "employment_type",
        "phone",
        "join_date",
        "termination_date",
        "accommodation_required",
        "shift_type",
        "notes",
    ):
        if field in data:
            setattr(emp, field, data[field])

    # BUG-007: flipping accommodation_required to false must NOT auto-release
    # an existing bed or delete the accommodation record. Checkout remains the
    # only release path (AGENTS.md business invariant 5 / user requirement).

    _assert_dates(emp.join_date, emp.termination_date)

    status_change = (
        routed_status is not None and routed_status != before.get("employment_status")
    )
    if routed_status is not None and not status_change:
        emp.employment_status = routed_status

    await audit_service.write_audit(
        db,
        action="employee.update",
        resource_type="employee",
        resource_id=emp.id,
        actor_user_id=actor_user_id,
        request_id=request_id,
        before_data=before,
        after_data=_to_out(emp),
    )
    if not status_change:
        await db.commit()
        return await get_employee(db, emp.id)

    # Persist profile edits, then apply the status event in a dedicated txn.
    term_date = emp.termination_date
    await db.commit()
    from app.services import offboarding_service

    result = await offboarding_service.apply_status_change(
        db,
        employee_id,
        new_status=routed_status,
        termination_date=term_date,
        actor_user_id=actor_user_id,
        request_id=request_id,
    )
    return await get_employee(db, result["employee_id"])


async def get_employee_accommodation(
    db: AsyncSession, employee_id: int, *, scope: DeptScope | None = None
) -> dict | None:
    """Current occupancy for an employee, or None when not housed."""
    emp = await _load(db, employee_id)
    if scope is not None:
        scope.require(emp.department_id, resource=f"employee {employee_id}")
    from app.models.accommodation import AccommodationRecord, BedReservation

    rec = (
        await db.execute(
            select(AccommodationRecord).where(
                AccommodationRecord.employee_id == employee_id,
                AccommodationRecord.status == "active",
            )
        )
    ).scalar_one_or_none()
    if rec is not None:
        from app.services import occupancy as occ

        chain = await occ.load_bed_chain(db, rec.bed_id)
        return {
            "kind": "accommodation",
            "id": rec.id,
            "bed_id": rec.bed_id,
            "checkin_at": rec.checkin_at,
            "status": rec.status,
            "bed_no": chain["bed"].bed_no,
            "room_no": chain["room"].room_no,
            "floor_no": chain["floor"].floor_no,
            "building_name": chain["building"].building_name,
            "area_name": chain["area"].area_name,
        }

    res = (
        await db.execute(
            select(BedReservation).where(
                BedReservation.employee_id == employee_id,
                BedReservation.status == "active",
            )
        )
    ).scalar_one_or_none()
    if res is not None:
        from app.services import occupancy as occ

        chain = await occ.load_bed_chain(db, res.bed_id)
        return {
            "kind": "reservation",
            "id": res.id,
            "bed_id": res.bed_id,
            "reserved_at": res.reserved_at,
            "expires_at": res.expires_at,
            "status": res.status,
            "bed_no": chain["bed"].bed_no,
            "room_no": chain["room"].room_no,
            "floor_no": chain["floor"].floor_no,
            "building_name": chain["building"].building_name,
            "area_name": chain["area"].area_name,
        }
    return None


async def get_employee_history(
    db: AsyncSession, employee_id: int, *, scope: DeptScope | None = None
) -> list[dict]:
    """Full accommodation history (all statuses, newest first)."""
    emp = await _load(db, employee_id)
    if scope is not None:
        scope.require(emp.department_id, resource=f"employee {employee_id}")
    from app.models.accommodation import AccommodationRecord
    from app.services import occupancy as occ

    rows = (
        await db.execute(
            select(AccommodationRecord)
            .where(AccommodationRecord.employee_id == employee_id)
            .order_by(AccommodationRecord.id.desc())
        )
    ).scalars().all()
    out: list[dict] = []
    for rec in rows:
        chain = await occ.load_bed_chain(db, rec.bed_id)
        out.append(
            {
                "id": rec.id,
                "bed_id": rec.bed_id,
                "checkin_at": rec.checkin_at,
                "checkout_at": rec.checkout_at,
                "status": rec.status,
                "end_reason": rec.end_reason,
                "bed_no": chain["bed"].bed_no,
                "room_no": chain["room"].room_no,
                "floor_no": chain["floor"].floor_no,
                "building_name": chain["building"].building_name,
                "area_name": chain["area"].area_name,
            }
        )
    return out


def _assert_dates(join_date: date | None, termination_date: date | None) -> None:
    if join_date and termination_date and termination_date < join_date:
        raise ValidationError("termination_date cannot be earlier than join_date")


async def _assert_department_active(db: AsyncSession, department_id: int) -> Department:
    dept = (
        await db.execute(select(Department).where(Department.id == department_id))
    ).scalar_one_or_none()
    if dept is None:
        raise ValidationError(f"Department {department_id} not found")
    if dept.status != "active":
        raise ValidationError(f"Department '{dept.department_code}' is disabled")
    return dept


# ---------------------------------------------------------------------------
# Excel template / import / export
# ---------------------------------------------------------------------------

IMPORT_COLUMNS = [
    ("employee_no", "员工工号", True),
    ("full_name", "员工姓名", True),
    ("gender", "性别", False),
    ("department_code", "部门编码", False),
    ("position", "岗位", False),
    ("join_date", "入职日期", False),
    ("phone", "联系电话", False),
    ("employment_type", "用工类型", False),
    ("shift_type", "班次", False),
    ("accommodation_required", "是否需要住宿", False),
]

_GENDER_MAP = {
    "男": "male",
    "女": "female",
    "其他": "other",
    "male": "male",
    "female": "female",
    "other": "other",
}
_EMP_TYPE_MAP = {
    "正式": "full_time",
    "全职": "full_time",
    "兼职": "part_time",
    "合同": "contract",
    "实习": "intern",
    "外包": "outsourced",
    "full_time": "full_time",
    "part_time": "part_time",
    "contract": "contract",
    "intern": "intern",
    "outsourced": "outsourced",
}
_SHIFT_MAP = {
    "白班": "day",
    "夜班": "night",
    "轮班": "rotating",
    "无": "none",
    "day": "day",
    "night": "night",
    "rotating": "rotating",
    "none": "none",
}


def build_employee_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "员工导入"
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    for col, (_key, label, _required) in enumerate(IMPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = 16
    sample = [
        "EMP001",
        "张三",
        "男",
        "DEPT001",
        "工程师",
        "2026-01-15",
        "13800000001",
        "正式",
        "白班",
        "是",
    ]
    for col, value in enumerate(sample, start=1):
        ws.cell(row=2, column=col, value=value)
    notes = wb.create_sheet("填写说明")
    notes["A1"] = "字段说明"
    notes["A1"].font = Font(bold=True)
    row = 2
    for key, label, required in IMPORT_COLUMNS:
        notes.cell(row=row, column=1, value=label)
        notes.cell(row=row, column=2, value=f"字段键 {key}；{'必填' if required else '可选'}")
        row += 1
    notes.cell(row=row + 1, column=1, value="导入以员工工号为唯一匹配键；日期格式 YYYY-MM-DD。")
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 48

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"invalid date '{text}', expected YYYY-MM-DD")


def _parse_bool(value: Any) -> bool:
    if value is None or value == "":
        return True
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"是", "y", "yes", "true", "1", "需要"}:
        return True
    if text in {"否", "n", "no", "false", "0", "不需要"}:
        return False
    raise ValueError(f"invalid boolean '{value}'")


def _map_enum(value: Any, mapping: dict[str, str], field: str) -> str | None:
    if value is None or value == "":
        return None
    text = str(value).strip()
    if text not in mapping:
        raise ValueError(f"invalid {field} '{text}'")
    return mapping[text]


def _file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _sanitize_import_text(value: Any) -> Any:
    return _safe_excel_value(value)


async def preview_employee_import(
    db: AsyncSession,
    *,
    filename: str,
    content: bytes,
    actor_user_id: int | None = None,
    request_id: str | None = None,
) -> dict:
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise ValidationError("Only .xlsx / .xlsm workbooks are supported")
    if len(content) > 5 * 1024 * 1024:
        raise ValidationError("File exceeds 5MB limit")
    if len(content) == 0:
        raise ValidationError("Uploaded file is empty")

    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surface as validation error
        raise ValidationError(f"Cannot parse workbook: {exc}") from exc

    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            raise ValidationError("Workbook has no rows") from None

        header_map: dict[str, int] = {}
        for idx, cell in enumerate(header):
            if cell is None:
                continue
            label = str(cell).strip()
            for key, col_label, _required in IMPORT_COLUMNS:
                if label == col_label or label == key:
                    header_map[key] = idx
        missing = [
            label
            for key, label, required in IMPORT_COLUMNS
            if required and key not in header_map
        ]
        if missing:
            raise ValidationError(f"Missing required columns: {', '.join(missing)}")

        dept_by_code: dict[str, Department] = {
            d.department_code: d for d in (await db.execute(select(Department))).scalars()
        }

        raw_rows: list[dict] = []
        seen_nos: set[str] = set()
        row_no = 1
        for values in rows_iter:
            row_no += 1
            if values is None or all(v is None or str(v).strip() == "" for v in values):
                continue
            if row_no > 2000 + 1:
                raise ValidationError("Import supports at most 2000 data rows")

            def cell(key: str, row_values: tuple = values) -> Any:
                idx = header_map.get(key)
                if idx is None or idx >= len(row_values):
                    return None
                return row_values[idx]

            result: dict[str, Any] = {
                "row_no": row_no,
                "action": "error",
                "employee_no": None,
                "full_name": None,
                "message": None,
                "data": None,
            }
            try:
                employee_no = str(cell("employee_no") or "").strip()
                full_name = str(cell("full_name") or "").strip()
                if not employee_no:
                    raise ValueError("employee_no is required")
                if not full_name:
                    raise ValueError("full_name is required")
                if employee_no in seen_nos:
                    raise ValueError(f"duplicate employee_no '{employee_no}' in file")
                seen_nos.add(employee_no)

                gender = _map_enum(cell("gender"), _GENDER_MAP, "gender")
                employment_type = _map_enum(
                    cell("employment_type"), _EMP_TYPE_MAP, "employment_type"
                )
                shift_type = _map_enum(cell("shift_type"), _SHIFT_MAP, "shift_type")
                join_date = _parse_date(cell("join_date"))
                phone = cell("phone")
                phone_text = (
                    str(phone).strip() if phone is not None and str(phone).strip() else None
                )
                if phone_text and not re.fullmatch(r"[0-9+\-\s]{5,32}", phone_text):
                    raise ValueError(f"invalid phone '{phone_text}'")
                department_code = (
                    str(cell("department_code")).strip()
                    if cell("department_code") is not None and str(cell("department_code")).strip()
                    else None
                )
                department_id = None
                if department_code:
                    dept = dept_by_code.get(department_code)
                    if dept is None:
                        raise ValueError(f"unknown department_code '{department_code}'")
                    department_id = dept.id

                accommodation_required = _parse_bool(cell("accommodation_required"))
                position = cell("position")
                position_text = (
                    str(position).strip()
                    if position is not None and str(position).strip()
                    else None
                )

                data = {
                    "employee_no": _sanitize_import_text(employee_no),
                    "full_name": _sanitize_import_text(full_name),
                    "gender": gender,
                    "department_id": department_id,
                    "department_code": department_code,
                    "position": _sanitize_import_text(position_text),
                    "employment_type": employment_type,
                    "phone": _sanitize_import_text(phone_text),
                    "join_date": join_date.isoformat() if join_date else None,
                    "accommodation_required": accommodation_required,
                    "shift_type": shift_type,
                }
                result["employee_no"] = employee_no
                result["full_name"] = full_name
                result["data"] = data

                existing = (
                    await db.execute(select(Employee).where(Employee.employee_no == employee_no))
                ).scalar_one_or_none()
                if existing is not None:
                    result["action"] = "update"
                    result["message"] = "employee_no exists, will update"
                else:
                    result["action"] = "create"
                    result["message"] = "new employee"
            except Exception as exc:  # noqa: BLE001 - collect per-row errors
                result["action"] = "error"
                result["message"] = str(exc)

            raw_rows.append(result)
    finally:
        # Always release the read-only workbook, including validation
        # error paths (missing headers, row-limit, empty sheet).
        wb.close()
    if not raw_rows:
        raise ValidationError("Workbook contains no data rows")

    create_count = sum(1 for r in raw_rows if r["action"] == "create")
    update_count = sum(1 for r in raw_rows if r["action"] == "update")
    error_count = sum(1 for r in raw_rows if r["action"] == "error")
    skip_count = sum(1 for r in raw_rows if r["action"] == "skip")

    job = ImportJob(
        import_type="employee",
        filename=_sanitize_import_text(filename)[:255],
        file_hash=_file_hash(content),
        status="previewed",
        total_rows=len(raw_rows),
        success_rows=create_count + update_count,
        failed_rows=error_count,
        created_by=actor_user_id,
        detail=json.dumps({"rows": raw_rows}, ensure_ascii=False, default=str),
    )
    db.add(job)
    await db.flush()
    await audit_service.write_audit(
        db,
        action="employee.import_preview",
        resource_type="import_job",
        resource_id=job.id,
        actor_user_id=actor_user_id,
        request_id=request_id,
        after_data={
            "filename": filename,
            "total_rows": len(raw_rows),
            "create_count": create_count,
            "update_count": update_count,
            "error_count": error_count,
        },
    )
    await db.commit()

    return {
        "import_job_id": job.id,
        "filename": filename,
        "total_rows": len(raw_rows),
        "create_count": create_count,
        "update_count": update_count,
        "skip_count": skip_count,
        "error_count": error_count,
        "rows": raw_rows,
    }


async def commit_employee_import(
    db: AsyncSession,
    import_job_id: int,
    *,
    actor_user_id: int | None = None,
    request_id: str | None = None,
) -> dict:
    job = (
        await db.execute(
            select(ImportJob)
            .where(ImportJob.id == import_job_id)
            .with_for_update()
            .execution_options(populate_existing=True)
        )
    ).scalar_one_or_none()
    if job is None:
        raise ResourceNotFound(f"Import job {import_job_id} not found")
    if job.import_type != "employee":
        raise ValidationError(f"Import job {import_job_id} is not an employee import")
    if job.status not in {"previewed", "failed"}:
        raise ConflictError(f"Import job {import_job_id} is already {job.status}")

    # Claim the job so a concurrent commit cannot pass the status check
    # (DMS-P2-005). Status is finalised below in the same transaction.
    job.status = "committing"

    if not job.detail:
        raise ValidationError("Import job has no preview snapshot")
    snapshot = json.loads(job.detail)
    rows: list[dict] = snapshot.get("rows", [])

    # Re-validate against current DB state so commit is not a blind replay.
    dept_by_id: dict[int, Department] = {
        d.id: d for d in (await db.execute(select(Department))).scalars()
    }
    created = 0
    updated = 0
    skipped = 0
    errors: list[dict] = []
    applied = 0

    for row in rows:
        action = row.get("action")
        data = row.get("data") or {}
        if action == "error":
            errors.append(row)
            continue
        if action == "skip":
            skipped += 1
            continue

        employee_no = data.get("employee_no")
        # Per-row SAVEPOINT: a DB failure on one row must not poison the
        # outer transaction or discard previously flushed good rows.
        try:
            async with db.begin_nested():
                if not employee_no:
                    raise ValueError("employee_no is required")
                department_id = data.get("department_id")
                if department_id is not None and department_id not in dept_by_id:
                    raise ValueError(f"department {department_id} no longer exists")
                if department_id is not None and dept_by_id[department_id].status != "active":
                    raise ValueError(
                        f"department '{dept_by_id[department_id].department_code}' is disabled"
                    )
                join_date = _parse_date(data.get("join_date"))

                existing = (
                    await db.execute(select(Employee).where(Employee.employee_no == employee_no))
                ).scalar_one_or_none()

                # Preview/commit intent must match. Drift is a conflict —
                # never silently flip create↔update (DMS-P1-002).
                if action == "create" and existing is not None:
                    raise ConflictError(
                        f"conflict: employee_no '{employee_no}' already exists "
                        f"(preview action=create); re-preview required"
                    )
                if action == "update" and existing is None:
                    raise ConflictError(
                        f"conflict: employee_no '{employee_no}' not found "
                        f"(preview action=update); re-preview required"
                    )

                fields = {
                    "full_name": data.get("full_name"),
                    "gender": data.get("gender"),
                    "department_id": department_id,
                    "position": data.get("position"),
                    "employment_type": data.get("employment_type"),
                    "phone": data.get("phone"),
                    "join_date": join_date,
                    "accommodation_required": bool(data.get("accommodation_required", True)),
                    "shift_type": data.get("shift_type"),
                }

                if action == "create":
                    emp = Employee(
                        employee_no=employee_no,
                        employment_status="active",
                        **fields,
                    )
                    db.add(emp)
                    await db.flush()
                    await audit_service.write_audit(
                        db,
                        action="employee.import_create",
                        resource_type="employee",
                        resource_id=emp.id,
                        actor_user_id=actor_user_id,
                        request_id=request_id,
                        after_data={"employee_no": employee_no, "import_job_id": job.id},
                    )
                else:
                    before = {
                        "employee_no": existing.employee_no,
                        "full_name": existing.full_name,
                        "department_id": existing.department_id,
                        "employment_status": existing.employment_status,
                    }
                    for key, value in fields.items():
                        setattr(existing, key, value)
                    await db.flush()
                    await audit_service.write_audit(
                        db,
                        action="employee.import_update",
                        resource_type="employee",
                        resource_id=existing.id,
                        actor_user_id=actor_user_id,
                        request_id=request_id,
                        before_data=before,
                        after_data={"employee_no": employee_no, "import_job_id": job.id},
                    )
            # Count only after the SAVEPOINT has been released successfully.
            if action == "create":
                created += 1
            else:
                updated += 1
            applied += 1
        except Exception as exc:  # noqa: BLE001 - collect per-row failures
            err_row = dict(row)
            err_row["action"] = "error"
            err_row["message"] = str(exc)
            errors.append(err_row)

    job.status = "completed" if not errors else "completed_with_errors"
    job.success_rows = applied
    job.failed_rows = len(errors)
    job.completed_at = utcnow()
    job.detail = json.dumps(
        {
            "summary": {
                "created": created,
                "updated": updated,
                "skipped": skipped,
                "errors": len(errors),
            },
            "errors": errors,
        },
        ensure_ascii=False,
        default=str,
    )

    await audit_service.write_audit(
        db,
        action="employee.import_commit",
        resource_type="import_job",
        resource_id=job.id,
        actor_user_id=actor_user_id,
        request_id=request_id,
        after_data={
            "created": created,
            "updated": updated,
            "skipped": skipped,
            "errors": len(errors),
        },
    )
    await db.commit()

    return {
        "import_job_id": job.id,
        "status": job.status,
        "total_rows": job.total_rows,
        "success_rows": job.success_rows,
        "failed_rows": job.failed_rows,
        "created_count": created,
        "updated_count": updated,
        "skipped_count": skipped,
        "errors": errors,
    }


async def get_import_job(db: AsyncSession, import_job_id: int) -> dict:
    job = (
        await db.execute(select(ImportJob).where(ImportJob.id == import_job_id))
    ).scalar_one_or_none()
    if job is None:
        raise ResourceNotFound(f"Import job {import_job_id} not found")
    return {
        "id": job.id,
        "import_type": job.import_type,
        "filename": job.filename,
        "file_hash": job.file_hash,
        "status": job.status,
        "total_rows": job.total_rows,
        "success_rows": job.success_rows,
        "failed_rows": job.failed_rows,
        "created_by": job.created_by,
        "completed_at": job.completed_at,
        "created_at": job.created_at,
    }


async def get_import_errors(db: AsyncSession, import_job_id: int) -> list[dict]:
    job = (
        await db.execute(select(ImportJob).where(ImportJob.id == import_job_id))
    ).scalar_one_or_none()
    if job is None:
        raise ResourceNotFound(f"Import job {import_job_id} not found")
    if not job.detail:
        return []
    snapshot = json.loads(job.detail)
    if "errors" in snapshot:
        return snapshot["errors"]
    rows = snapshot.get("rows", [])
    return [r for r in rows if r.get("action") == "error"]


async def export_employees(
    db: AsyncSession,
    *,
    keyword: str | None = None,
    department_id: int | None = None,
    employment_status: str | None = None,
    scope: DeptScope | None = None,
) -> bytes:
    from app.services.data_scope import apply_employee_scope

    query = select(Employee)
    if scope is not None:
        query = apply_employee_scope(query, scope, Employee.id)
    if keyword:
        pattern = f"%{keyword}%"
        query = query.where(
            (Employee.employee_no.ilike(pattern)) | (Employee.full_name.ilike(pattern))
        )
    if department_id is not None:
        query = query.where(Employee.department_id == department_id)
    if employment_status:
        query = query.where(Employee.employment_status == employment_status)

    rows = (await db.execute(query.order_by(Employee.employee_no))).scalars().all()
    depts = await _dept_map(db, {r.department_id for r in rows if r.department_id})

    wb = Workbook()
    ws = wb.active
    ws.title = "员工导出"
    headers = [
        "员工工号",
        "员工姓名",
        "性别",
        "部门编码",
        "部门名称",
        "岗位",
        "入职日期",
        "联系电话",
        "用工类型",
        "班次",
        "人事状态",
        "是否需要住宿",
        "住宿状态",
        "备注",
    ]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True)

    gender_rev = {"male": "男", "female": "女", "other": "其他"}
    type_rev = {
        "full_time": "正式",
        "part_time": "兼职",
        "contract": "合同",
        "intern": "实习",
        "outsourced": "外包",
    }
    shift_rev = {"day": "白班", "night": "夜班", "rotating": "轮班", "none": "无"}
    status_rev = {
        "pending": "待入职",
        "active": "在职",
        "probation": "试用期",
        "offboarding": "离职办理中",
        "terminated": "已离职",
        "resigned": "已离职",
        "suspended": "停职",
    }

    for i, emp in enumerate(rows, start=2):
        dept = depts.get(emp.department_id) if emp.department_id else None
        values = [
            _safe_excel_value(emp.employee_no),
            _safe_excel_value(emp.full_name),
            gender_rev.get(emp.gender or "", emp.gender),
            dept.department_code if dept else None,
            dept.department_name if dept else None,
            _safe_excel_value(emp.position),
            emp.join_date.isoformat() if emp.join_date else None,
            _safe_excel_value(emp.phone),
            type_rev.get(emp.employment_type or "", emp.employment_type),
            shift_rev.get(emp.shift_type or "", emp.shift_type),
            status_rev.get(emp.employment_status, emp.employment_status),
            "是" if emp.accommodation_required else "否",
            "未入住",
            _safe_excel_value(emp.notes),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=i, column=col, value=value)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
