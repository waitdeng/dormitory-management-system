"""V1.0.1 BUG-005 + UX-006 — bed full path and Chinese status labels (real MySQL)."""

from __future__ import annotations

import io

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from tests.helpers_m2 import admin_headers
from tests.helpers_m3 import make_dorm, make_employee


@pytest.mark.asyncio
async def test_bug005_available_beds_carry_full_path(client: AsyncClient):
    admin = await admin_headers(client)
    dorm = await make_dorm(client, admin, capacity=1)
    resp = await client.get(
        "/api/v1/beds/available",
        headers=admin,
        params={"limit": 200, "building_id": dorm["building"]["id"]},
    )
    assert resp.status_code == 200, resp.text
    beds = resp.json()
    target = next(b for b in beds if b["id"] == dorm["beds"][0]["id"])
    assert target["area_name"] == dorm["area"]["area_name"]
    assert target["building_name"] == dorm["building"]["building_name"]
    assert target["floor_no"] == dorm["floor"]["floor_no"]
    assert target["room_no"] == dorm["room"]["room_no"]
    assert target["bed_no"] == dorm["beds"][0]["bed_no"]
    assert target["bed_path"]
    assert dorm["area"]["area_name"] in target["bed_path"]
    assert dorm["beds"][0]["bed_no"] in target["bed_path"]


@pytest.mark.asyncio
async def test_bug005_transfer_detail_beds_have_full_path(client: AsyncClient):
    admin = await admin_headers(client)
    dorm = await make_dorm(client, admin, capacity=2)
    emp = await make_employee(client, admin, suffix="路径")
    from tests.helpers_m4 import checkin_employee

    await checkin_employee(client, admin, emp["id"], dorm["beds"][0]["id"])
    tr = await client.post(
        "/api/v1/transfers",
        headers=admin,
        json={
            "employee_id": emp["id"],
            "target_bed_id": dorm["beds"][1]["id"],
            "reason": "path check",
        },
    )
    assert tr.status_code == 201, tr.text
    detail = await client.get(f"/api/v1/transfers/{tr.json()['id']}", headers=admin)
    assert detail.status_code == 200, detail.text
    body = detail.json()
    for key in ("source_bed", "target_bed"):
        bed = body[key]
        assert bed.get("area_name"), f"{key} missing area_name"
        assert bed.get("building_name"), f"{key} missing building_name"
        assert bed.get("floor_no") is not None, f"{key} missing floor_no"
        assert bed.get("room_no"), f"{key} missing room_no"
        assert bed.get("bed_no"), f"{key} missing bed_no"


@pytest.mark.asyncio
async def test_ux006_employee_export_uses_chinese_status(client: AsyncClient):
    admin = await admin_headers(client)
    await make_employee(client, admin, suffix="在职导出")
    resp = await client.get("/api/v1/employees/export", headers=admin)
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    idx = header.index("人事状态") if "人事状态" in header else None
    assert idx is not None, header
    statuses = {r[idx] for r in rows[1:] if r[idx]}
    assert statuses  # at least one
    for s in statuses:
        assert s not in {
            "active",
            "pending",
            "probation",
            "offboarding",
            "terminated",
            "resigned",
            "suspended",
        }, f"raw status leaked: {s}"
    assert any(s in {"在职", "试用期", "待入职", "已离职", "离职办理中", "停职"} for s in statuses)


@pytest.mark.asyncio
async def test_ux006_report_export_uses_chinese_status(client: AsyncClient):
    admin = await admin_headers(client)
    dorm = await make_dorm(client, admin, capacity=1)
    emp = await make_employee(client, admin, suffix="台账导出")
    from tests.helpers_m4 import checkin_employee

    await checkin_employee(client, admin, emp["id"], dorm["beds"][0]["id"])
    resp = await client.get("/api/v1/reports/accommodations/export", headers=admin)
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    idx = header.index("人事状态")
    statuses = {r[idx] for r in rows[1:] if r[idx]}
    assert "在职" in statuses
    assert "active" not in statuses
