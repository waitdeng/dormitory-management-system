"""M6 historical employee Excel import drill (live API).

Generates a multi-hundred-row historical workbook (including dirty rows and a
formula-injection attempt), runs the two-phase import (preview → commit), and
reports create/update/error counts. Does not touch dorm assignment history —
V1.0 employee import is the supported historical onboarding path.

Usage:
  cd backend && .venv/bin/python tests/historical_import_drill.py
Env: PERF_BASE_URL / DMS_ADMIN_USERNAME / DMS_ADMIN_PASSWORD
"""

from __future__ import annotations

import os
import sys
import uuid
from io import BytesIO

import httpx
from openpyxl import Workbook

BASE = os.environ.get("PERF_BASE_URL", "http://127.0.0.1:8080")
ADMIN = os.environ.get("DMS_ADMIN_USERNAME", "admin")
PASSWORD = os.environ.get("DMS_ADMIN_PASSWORD", "")
ROW_COUNT = int(os.environ.get("HIST_IMPORT_ROWS", "300"))

HEADERS = [
    "员工工号",
    "员工姓名",
    "性别",
    "部门编码",
    "岗位",
    "入职日期",
    "联系电话",
    "用工类型",
    "班次",
    "是否需要住宿",
]


def build_workbook(dept_code: str, batch: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "员工导入"
    ws.append(HEADERS)
    for i in range(ROW_COUNT):
        no = f"H{batch}{i:04d}"
        gender = "男" if i % 2 == 0 else "女"
        year = 2015 + (i % 8)
        month = 1 + (i % 12)
        day = 1 + (i % 28)
        hire = f"{year:04d}-{month:02d}-{day:02d}"
        phone = f"139{(i % 100000000):08d}"
        ws.append(
            [
                no,
                f"历史员工{i:04d}",
                gender,
                dept_code,
                "专员" if i % 3 else "主管",
                hire,
                phone,
                "正式",
                "白班",
                "是" if i % 4 else "否",
            ]
        )
    # dirty / edge rows
    ws.append(["", "空工号", "男", dept_code, "", "2020-01-01", "", "", "", ""])
    ws.append([f"H{batch}DUP", "重复A", "男", dept_code, "", "2020-01-02", "", "", "", ""])
    ws.append([f"H{batch}DUP", "重复B", "女", dept_code, "", "2020-01-03", "", "", "", ""])
    ws.append(
        [f"H{batch}INJ", "=cmd|'/c calc'!A1", "男", dept_code, "", "2020-01-04", "", "", "", ""]
    )
    ws.append([f"H{batch}BAD", "坏部门", "男", "NO_SUCH_DEPT", "", "not-a-date", "", "", "", ""])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> int:
    if not PASSWORD:
        print("Set DMS_ADMIN_PASSWORD", file=sys.stderr)
        return 2

    batch = uuid.uuid4().hex[:6].upper()
    with httpx.Client(base_url=BASE, timeout=120.0, trust_env=False) as client:
        login = client.post(
            "/api/v1/auth/login", json={"username": ADMIN, "password": PASSWORD}
        )
        login.raise_for_status()
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

        dept_code = f"HDEPT{batch}"
        r = client.post(
            "/api/v1/departments",
            headers=headers,
            json={"department_code": dept_code, "department_name": f"历史导入部门{batch}"},
        )
        if r.status_code not in (200, 201):
            print("create department failed", r.status_code, r.text)
            return 1

        content = build_workbook(dept_code, batch)
        xlsx_ct = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        preview = client.post(
            "/api/v1/employees/import/preview",
            headers=headers,
            files={"file": ("historical.xlsx", content, xlsx_ct)},
        )
        if preview.status_code != 200:
            print("preview failed", preview.status_code, preview.text[:500])
            return 1
        body = preview.json()
        job_id = body["import_job_id"]
        print("PREVIEW", {k: body.get(k) for k in ("total_rows", "create_count", "update_count", "error_count", "skip_count")})

        commit = client.post(
            "/api/v1/employees/import/commit",
            headers=headers,
            params={"import_job_id": job_id},
        )
        if commit.status_code != 200:
            print("commit failed", commit.status_code, commit.text[:500])
            return 1
        result = commit.json()
        print("COMMIT", {k: result.get(k) for k in ("created_count", "updated_count", "failed_rows", "skipped_count")})

        job = client.get(f"/api/v1/imports/{job_id}", headers=headers)
        print("JOB", job.json().get("status"))
        errors = client.get(f"/api/v1/imports/{job_id}/errors", headers=headers).json()
        print("ERROR_ROWS", len(errors) if isinstance(errors, list) else errors)

        # verify a sample created employee is queryable
        sample_no = f"H{batch}0000"
        found = client.get(
            "/api/v1/employees", headers=headers, params={"keyword": sample_no}
        ).json()
        ok = found.get("total", 0) >= 1
        print("SAMPLE_LOOKUP", sample_no, "ok" if ok else "MISSING")

        # export must not contain raw formula cells
        export = client.get("/api/v1/employees/export", headers=headers)
        print("EXPORT", export.status_code, "bytes", len(export.content))
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(export.content))
        formulas = [
            c
            for row in wb.active.iter_rows(values_only=True)
            for c in row
            if isinstance(c, str) and c.startswith("=")
        ]
        print("EXPORT_FORMULAS", len(formulas))

        if not ok or formulas:
            return 1
        print("HISTORICAL_IMPORT_DRILL PASS")
        return 0


if __name__ == "__main__":
    raise SystemExit(main())
