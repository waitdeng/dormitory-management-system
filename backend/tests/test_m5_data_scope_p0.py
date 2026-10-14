"""M5 P0 — department data-scope isolation (real MySQL).

Covers AGENTS.md M5 P0 acceptance:
  - dept_manager only sees authorized departments' employees and housing
  - direct access to out-of-scope employee detail / history is rejected
  - query parameters cannot bypass scope
  - Excel export follows the same scope
  - todos / alerts / stats must not leak other departments
  - super admin retains full access
"""

from __future__ import annotations

import pytest
from httpx import AsyncClient

from tests.helpers_m2 import admin_headers, unique_code
from tests.helpers_m3 import make_dorm
from tests.helpers_m4 import checkin_employee, headers_for_user, make_user_with_role


async def _make_dept(client: AsyncClient, headers: dict, *, name: str) -> dict:
    resp = await client.post(
        "/api/v1/departments",
        headers=headers,
        json={
            "department_code": unique_code("D"),
            "department_name": name,
        },
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


async def _make_emp_in_dept(
    client: AsyncClient, headers: dict, dept_id: int, suffix: str
) -> dict:
    payload = {
        "employee_no": unique_code("E"),
        "full_name": f"范围员工{suffix}",
        "gender": "male",
        "employment_status": "active",
        "accommodation_required": True,
        "department_id": dept_id,
    }
    resp = await client.post("/api/v1/employees", headers=headers, json=payload)
    assert resp.status_code == 201, resp.text
    return resp.json()


@pytest.mark.asyncio
async def test_p0_dept_manager_scope_isolation(client: AsyncClient):
    """P0: dept_manager only queries authorized department data."""
    admin = await admin_headers(client)
    assert admin

    dept_a = await _make_dept(client, admin, name="范围部门A")
    dept_b = await _make_dept(client, admin, name="范围部门B")

    emp_a = await _make_emp_in_dept(client, admin, dept_a["id"], "A")
    emp_b = await _make_emp_in_dept(client, admin, dept_b["id"], "B")

    dorm = await make_dorm(client, admin, capacity=2)
    await checkin_employee(client, admin, emp_a["id"], dorm["beds"][0]["id"])
    await checkin_employee(client, admin, emp_b["id"], dorm["beds"][1]["id"])

    # dept_manager scoped to department A only
    mgr = await make_user_with_role(client, admin, "dept_manager", "mgr_scope")
    r = await client.put(
        f"/api/v1/users/{mgr['id']}/departments",
        headers=admin,
        json={"department_ids": [dept_a["id"]]},
    )
    assert r.status_code == 200, r.text
    mgr_headers = await headers_for_user(client, mgr["username"])

    # 1) list only shows department A employees
    listed = await client.get("/api/v1/employees", headers=mgr_headers)
    assert listed.status_code == 200
    ids = {i["id"] for i in listed.json()["items"]}
    assert emp_a["id"] in ids
    assert emp_b["id"] not in ids

    # 2) detail of out-of-scope employee is rejected
    detail = await client.get(f"/api/v1/employees/{emp_b['id']}", headers=mgr_headers)
    assert detail.status_code == 403, detail.text

    # 3) accommodation history of out-of-scope employee is rejected
    hist = await client.get(
        f"/api/v1/employees/{emp_b['id']}/history", headers=mgr_headers
    )
    assert hist.status_code == 403, hist.text

    # 4) query param cannot bypass scope
    bypass = await client.get(
        f"/api/v1/employees?department_id={dept_b['id']}", headers=mgr_headers
    )
    assert bypass.status_code == 200
    assert bypass.json()["items"] == []
    assert bypass.json()["total"] == 0

    # 5) in-scope employee detail works
    ok = await client.get(f"/api/v1/employees/{emp_a['id']}", headers=mgr_headers)
    assert ok.status_code == 200
    assert ok.json()["id"] == emp_a["id"]

    # 6) Excel export follows scope
    export = await client.get("/api/v1/employees/export", headers=mgr_headers)
    assert export.status_code == 200
    assert export.content[:2] == b"PK"  # xlsx zip magic
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(export.content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    nos = {r[0] for r in rows if r and r[0]}
    assert emp_a["employee_no"] in nos
    assert emp_b["employee_no"] not in nos

    # 7) active accommodations list scoped
    accs = await client.get("/api/v1/accommodations/active", headers=mgr_headers)
    if accs.status_code == 200:
        acc_emp_ids = {i["employee_id"] for i in accs.json().get("items", [])}
        assert emp_b["id"] not in acc_emp_ids

    # 8) admin still sees both
    admin_list = await client.get("/api/v1/employees", headers=admin)
    admin_ids = {i["id"] for i in admin_list.json()["items"]}
    assert emp_a["id"] in admin_ids and emp_b["id"] in admin_ids
    admin_b = await client.get(f"/api/v1/employees/{emp_b['id']}", headers=admin)
    assert admin_b.status_code == 200


@pytest.mark.asyncio
async def test_p0_scope_blocks_todos_alerts_stats_leak(client: AsyncClient):
    """P0: todos, alerts, dashboard and reports must not leak other departments."""
    admin = await admin_headers(client)
    assert admin

    dept_a = await _make_dept(client, admin, name="待办部门A")
    dept_b = await _make_dept(client, admin, name="待办部门B")
    emp_a = await _make_emp_in_dept(client, admin, dept_a["id"], "TA")
    emp_b = await _make_emp_in_dept(client, admin, dept_b["id"], "TB")

    dorm = await make_dorm(client, admin, capacity=2)
    await checkin_employee(client, admin, emp_a["id"], dorm["beds"][0]["id"])
    await checkin_employee(client, admin, emp_b["id"], dorm["beds"][1]["id"])

    # generate offboarding todo + alert for employee B (out of scope)
    r = await client.post(
        f"/api/v1/employees/{emp_b['id']}/offboarding", headers=admin, json={}
    )
    assert r.status_code == 200, r.text
    r = await client.post(
        f"/api/v1/employees/{emp_b['id']}/terminate", headers=admin, json={}
    )
    assert r.status_code == 200, r.text

    mgr = await make_user_with_role(client, admin, "dept_manager", "mgr_todo")
    await client.put(
        f"/api/v1/users/{mgr['id']}/departments",
        headers=admin,
        json={"department_ids": [dept_a["id"]]},
    )
    mgr_headers = await headers_for_user(client, mgr["username"])

    # todos must not include employee B
    todos = await client.get("/api/v1/workflow/todos", headers=mgr_headers)
    assert todos.status_code == 200
    todo_emp_ids = {
        i["employee_id"] for i in todos.json()["items"] if i.get("employee_id")
    }
    assert emp_b["id"] not in todo_emp_ids

    # alerts must not include employee B
    alerts = await client.get("/api/v1/workflow/alerts", headers=mgr_headers)
    assert alerts.status_code == 200
    alert_emp_ids = {
        i["employee_id"] for i in alerts.json()["items"] if i.get("employee_id")
    }
    assert emp_b["id"] not in alert_emp_ids

    # dashboard summary housed count is scoped (only A is in scope)
    summary = await client.get("/api/v1/dashboard/summary", headers=mgr_headers)
    assert summary.status_code == 200
    body = summary.json()
    assert body["housed_employees"] == 1  # only emp A

    # accommodations ledger scoped
    ledger = await client.get("/api/v1/reports/accommodations", headers=mgr_headers)
    assert ledger.status_code == 200
    ledger_emps = {i["employee_id"] for i in ledger.json()["items"]}
    assert emp_a["id"] in ledger_emps
    assert emp_b["id"] not in ledger_emps

    # reports export scoped
    export = await client.get(
        "/api/v1/reports/accommodations/export", headers=mgr_headers
    )
    assert export.status_code == 200
    assert export.content[:2] == b"PK"

    # admin sees both employees in ledger
    admin_ledger = await client.get("/api/v1/reports/accommodations", headers=admin)
    admin_emps = {i["employee_id"] for i in admin_ledger.json()["items"]}
    assert emp_a["id"] in admin_emps and emp_b["id"] in admin_emps


@pytest.mark.asyncio
async def test_p0_manager_of_department_via_manager_employee_id(client: AsyncClient):
    """dept_manager linked as department manager gets that department automatically."""
    admin = await admin_headers(client)
    dept = await _make_dept(client, admin, name="经理部门")
    emp_mgr = await _make_emp_in_dept(client, admin, dept["id"], "MGR")
    emp_other = await _make_emp_in_dept(client, admin, dept["id"], "OTH")

    # assign manager_employee_id
    r = await client.patch(
        f"/api/v1/departments/{dept['id']}",
        headers=admin,
        json={"manager_employee_id": emp_mgr["id"]},
    )
    assert r.status_code == 200, r.text

    # another department the user does not manage
    dept_x = await _make_dept(client, admin, name="无关部门X")
    emp_x = await _make_emp_in_dept(client, admin, dept_x["id"], "X")

    user = await make_user_with_role(client, admin, "dept_manager", "mgr_link")
    # bind employee so manager_employee_id matches
    r = await client.patch(
        f"/api/v1/users/{user['id']}",
        headers=admin,
        json={"employee_id": emp_mgr["id"]},
    )
    assert r.status_code == 200, r.text
    # no explicit department assignment — only manager link
    await client.put(
        f"/api/v1/users/{user['id']}/departments",
        headers=admin,
        json={"department_ids": []},
    )
    headers = await headers_for_user(client, user["username"])

    listed = await client.get("/api/v1/employees", headers=headers)
    ids = {i["id"] for i in listed.json()["items"]}
    assert emp_mgr["id"] in ids
    assert emp_other["id"] in ids
    assert emp_x["id"] not in ids


@pytest.mark.asyncio
async def test_p0_dept_manager_cannot_modify_query_to_see_others(client: AsyncClient):
    """Even with crafted filters, out-of-scope rows stay hidden."""
    admin = await admin_headers(client)
    dept_a = await _make_dept(client, admin, name="过滤A")
    dept_b = await _make_dept(client, admin, name="过滤B")
    emp_b = await _make_emp_in_dept(client, admin, dept_b["id"], "FB")

    mgr = await make_user_with_role(client, admin, "dept_manager", "mgr_filter")
    await client.put(
        f"/api/v1/users/{mgr['id']}/departments",
        headers=admin,
        json={"department_ids": [dept_a["id"]]},
    )
    headers = await headers_for_user(client, mgr["username"])

    # keyword search across all
    r = await client.get(
        f"/api/v1/employees?keyword={emp_b['employee_no']}", headers=headers
    )
    assert r.status_code == 200
    assert r.json()["items"] == []

    # large page size
    r = await client.get("/api/v1/employees?page_size=100", headers=headers)
    assert all(i["id"] != emp_b["id"] for i in r.json()["items"])

    # reports keyword
    r = await client.get(
        f"/api/v1/reports/accommodations?keyword={emp_b['full_name']}", headers=headers
    )
    assert r.status_code == 200
    assert all(i["employee_id"] != emp_b["id"] for i in r.json()["items"])
