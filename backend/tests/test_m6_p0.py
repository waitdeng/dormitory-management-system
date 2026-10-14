"""M6 P0 acceptance tests: department data isolation on dashboard metrics."""

from __future__ import annotations

import pytest
from httpx import AsyncClient

from tests.helpers_m2 import admin_headers
from tests.helpers_m3 import make_dorm
from tests.helpers_m4 import (
    checkin_employee,
    headers_for_user,
    make_user_with_role,
)
from tests.test_m5_data_scope_p0 import _make_dept, _make_emp_in_dept


@pytest.mark.asyncio
async def test_p0_dashboard_employee_metrics_respect_dept_scope(client: AsyncClient):
    """P0-1: dashboard employee-linked KPIs must not leak out-of-scope departments.

    Facility capacity (bed_total / vacant_beds) stays global by design (D-036);
    employee activity and workflow counters are scoped.
    """
    admin = await admin_headers(client)
    assert admin

    dept_a = await _make_dept(client, admin, name="M6部门A")
    dept_b = await _make_dept(client, admin, name="M6部门B")
    emp_a = await _make_emp_in_dept(client, admin, dept_a["id"], "M6A")
    emp_b = await _make_emp_in_dept(client, admin, dept_b["id"], "M6B")

    dorm = await make_dorm(client, admin, capacity=2)
    await checkin_employee(client, admin, emp_a["id"], dorm["beds"][0]["id"])
    await checkin_employee(client, admin, emp_b["id"], dorm["beds"][1]["id"])

    # out-of-scope employee offboarding + terminate → todo + alert for B only
    r = await client.post(
        f"/api/v1/employees/{emp_b['id']}/offboarding", headers=admin, json={}
    )
    assert r.status_code == 200, r.text
    r = await client.post(
        f"/api/v1/employees/{emp_b['id']}/terminate", headers=admin, json={}
    )
    assert r.status_code == 200, r.text

    mgr = await make_user_with_role(client, admin, "dept_manager", "mgr_m6_dash")
    await client.put(
        f"/api/v1/users/{mgr['id']}/departments",
        headers=admin,
        json={"department_ids": [dept_a["id"]]},
    )
    mgr_headers = await headers_for_user(client, mgr["username"])

    summary = await client.get("/api/v1/dashboard/summary", headers=mgr_headers)
    assert summary.status_code == 200
    body = summary.json()

    # employee-linked metrics scoped to dept A only
    assert body["housed_employees"] == 1
    assert body["today_checkin"] == 1
    assert body["terminated_still_housed"] == 0
    assert body["active_alerts"] == 0
    assert body["pending_todos"] == 0

    # department_housed must not name or count dept B
    dept_ids = {d["department_id"] for d in body["department_housed"]}
    assert dept_a["id"] in dept_ids
    assert dept_b["id"] not in dept_ids
    for d in body["department_housed"]:
        assert d["department_id"] != dept_b["id"]
        if d["department_id"] == dept_a["id"]:
            assert d["housed_count"] == 1

    # facility capacity remains visible (allocation needs vacant inventory)
    assert body["bed_total"] >= 2
    assert body["vacant_beds"] >= 0

    # admin still sees both sides (department chart is capped at top 50 by housed)
    admin_summary = await client.get("/api/v1/dashboard/summary", headers=admin)
    assert admin_summary.status_code == 200
    ab = admin_summary.json()
    assert ab["housed_employees"] >= 2
    assert ab["terminated_still_housed"] >= 1
    assert ab["pending_todos"] >= 1
    assert ab["today_checkin"] >= 2


@pytest.mark.asyncio
async def test_p0_dashboard_trends_scoped(client: AsyncClient):
    """P0-1: check-in/out trend series must follow department scope."""
    admin = await admin_headers(client)
    dept_a = await _make_dept(client, admin, name="M6趋势A")
    dept_b = await _make_dept(client, admin, name="M6趋势B")
    emp_a = await _make_emp_in_dept(client, admin, dept_a["id"], "M6TA")
    emp_b = await _make_emp_in_dept(client, admin, dept_b["id"], "M6TB")

    dorm = await make_dorm(client, admin, capacity=2)
    await checkin_employee(client, admin, emp_a["id"], dorm["beds"][0]["id"])
    await checkin_employee(client, admin, emp_b["id"], dorm["beds"][1]["id"])

    mgr = await make_user_with_role(client, admin, "dept_manager", "mgr_m6_trend")
    await client.put(
        f"/api/v1/users/{mgr['id']}/departments",
        headers=admin,
        json={"department_ids": [dept_a["id"]]},
    )
    mgr_headers = await headers_for_user(client, mgr["username"])

    trends = await client.get("/api/v1/dashboard/trends", headers=mgr_headers)
    assert trends.status_code == 200
    items = trends.json()["items"]
    assert items
    total_checkin = sum(i["checkin"] for i in items)
    # only emp_a is in scope
    assert total_checkin == 1

    admin_trends = await client.get("/api/v1/dashboard/trends", headers=admin)
    admin_checkin = sum(i["checkin"] for i in admin_trends.json()["items"])
    assert admin_checkin >= 2


@pytest.mark.asyncio
async def test_p0_empty_scope_sees_no_employee_dashboard_data(client: AsyncClient):
    """P0-1: user with no authorized departments gets zero employee-linked KPIs."""
    admin = await admin_headers(client)
    dept = await _make_dept(client, admin, name="M6空范围")
    emp = await _make_emp_in_dept(client, admin, dept["id"], "M6EMP")
    dorm = await make_dorm(client, admin, capacity=1)
    await checkin_employee(client, admin, emp["id"], dorm["beds"][0]["id"])

    viewer = await make_user_with_role(client, admin, "dept_manager", "mgr_m6_empty")
    # no PUT /departments → empty scope
    headers = await headers_for_user(client, viewer["username"])
    summary = await client.get("/api/v1/dashboard/summary", headers=headers)
    assert summary.status_code == 200
    body = summary.json()
    assert body["housed_employees"] == 0
    assert body["today_checkin"] == 0
    assert body["terminated_still_housed"] == 0
    assert body["active_alerts"] == 0
    assert body["pending_todos"] == 0
    assert body["department_housed"] == []


@pytest.mark.asyncio
async def test_p0_full_lifecycle_hire_to_archive(client: AsyncClient):
    """P0-3: hire → allocate → check-in → transfer → offboard → alert → checkout → archive."""
    admin = await admin_headers(client)
    assert admin

    # 1. 入职：创建员工
    from tests.helpers_m3 import make_employee

    emp = await make_employee(client, admin, suffix="全周期")
    assert emp["employment_status"] in ("active", "probation", "pending")
    assert emp["accommodation_status"] == "not_in_dorm"

    # 2. 分配宿舍并正式入住
    dorm = await make_dorm(client, admin, capacity=2)
    bed1, bed2 = dorm["beds"][0]["id"], dorm["beds"][1]["id"]
    await checkin_employee(client, admin, emp["id"], bed1)

    acc = (await client.get(f"/api/v1/employees/{emp['id']}/accommodation", headers=admin)).json()
    active = acc if isinstance(acc, list) else acc.get("items", acc)
    if isinstance(active, dict):
        active = [active]
    assert len(active) == 1
    assert active[0]["status"] == "active"
    assert active[0]["bed_id"] == bed1

    emp_after_in = (await client.get(f"/api/v1/employees/{emp['id']}", headers=admin)).json()
    assert emp_after_in["accommodation_status"] == "in_dorm"

    beds = (await client.get(f"/api/v1/rooms/{dorm['room']['id']}/beds", headers=admin)).json()
    bed_list = beds["items"] if isinstance(beds, dict) else beds
    assert next(b for b in bed_list if b["id"] == bed1)["available"] is False

    # 3. 调宿：bed1 → bed2，历史保留
    tr = await client.post(
        "/api/v1/transfers",
        headers=admin,
        json={"employee_id": emp["id"], "target_bed_id": bed2, "reason": "M6调宿"},
    )
    assert tr.status_code == 201, tr.text
    tr_id = tr.json()["id"]
    r = await client.post(f"/api/v1/transfers/{tr_id}/reserve", headers=admin, json={})
    assert r.status_code == 200, r.text
    r = await client.post(f"/api/v1/transfers/{tr_id}/complete", headers=admin, json={})
    assert r.status_code == 200, r.text

    history = (await client.get(f"/api/v1/employees/{emp['id']}/history", headers=admin)).json()
    hist = history if isinstance(history, list) else history.get("items", [])
    assert len(hist) == 2
    ended = next(h for h in hist if h["status"] == "ended")
    act = next(h for h in hist if h["status"] == "active")
    assert ended["bed_id"] == bed1 and ended["end_reason"] == "transfer"
    assert act["bed_id"] == bed2

    beds = (await client.get(f"/api/v1/rooms/{dorm['room']['id']}/beds", headers=admin)).json()
    bed_list = beds["items"] if isinstance(beds, dict) else beds
    assert next(b for b in bed_list if b["id"] == bed1)["available"] is True
    assert next(b for b in bed_list if b["id"] == bed2)["available"] is False

    # 4. 离职办理 → 退宿待办 + 异常预警，床位不自动释放
    r = await client.post(f"/api/v1/employees/{emp['id']}/offboarding", headers=admin, json={})
    assert r.status_code == 200, r.text
    assert r.json()["side_effects"].get("todo_id")

    r = await client.post(f"/api/v1/employees/{emp['id']}/terminate", headers=admin, json={})
    assert r.status_code == 200, r.text

    todos = (await client.get("/api/v1/workflow/todos", headers=admin, params={"status": "pending"})).json()
    life_todos = [t for t in todos["items"] if t.get("employee_id") == emp["id"]]
    assert any(t["todo_type"] == "checkout_offboarding" for t in life_todos)

    alerts = (await client.get("/api/v1/workflow/alerts", headers=admin, params={"status": "active"})).json()
    life_alerts = [a for a in alerts["items"] if a.get("employee_id") == emp["id"]]
    assert any(a["alert_type"] == "terminated_still_housed" for a in life_alerts)

    emp_term = (await client.get(f"/api/v1/employees/{emp['id']}", headers=admin)).json()
    assert emp_term["employment_status"] == "terminated"
    assert emp_term["accommodation_status"] == "in_dorm"  # 离职不自动释放床位

    beds = (await client.get(f"/api/v1/rooms/{dorm['room']['id']}/beds", headers=admin)).json()
    bed_list = beds["items"] if isinstance(beds, dict) else beds
    assert next(b for b in bed_list if b["id"] == bed2)["available"] is False

    # 5. 行政确认退宿 → 释放床位、关闭待办/预警、历史归档
    cos = (await client.get("/api/v1/checkouts", headers=admin, params={"employee_id": emp["id"]})).json()
    open_cos = [c for c in cos["items"] if c["status"] in ("pending", "inspection")]
    assert len(open_cos) == 1
    co_id = open_cos[0]["id"]

    r = await client.patch(
        f"/api/v1/checkouts/{co_id}/inspection",
        headers=admin,
        json={
            "key_returned": True,
            "access_card_returned": True,
            "belongings_cleared": True,
            "room_inspected": True,
            "inspection_notes": "M6全周期检查完成",
        },
    )
    assert r.status_code == 200, r.text
    r = await client.post(f"/api/v1/checkouts/{co_id}/complete", headers=admin, json={})
    assert r.status_code == 200, r.text

    emp_done = (await client.get(f"/api/v1/employees/{emp['id']}", headers=admin)).json()
    assert emp_done["accommodation_status"] == "not_in_dorm"

    history = (await client.get(f"/api/v1/employees/{emp['id']}/history", headers=admin)).json()
    hist = history if isinstance(history, list) else history.get("items", [])
    assert len(hist) == 2
    assert all(h["status"] == "ended" for h in hist)
    assert all(h["checkout_at"] is not None for h in hist)

    beds = (await client.get(f"/api/v1/rooms/{dorm['room']['id']}/beds", headers=admin)).json()
    bed_list = beds["items"] if isinstance(beds, dict) else beds
    assert next(b for b in bed_list if b["id"] == bed2)["available"] is True

    todos = (
        await client.get(
            "/api/v1/workflow/todos", headers=admin, params={"employee_id": emp["id"]}
        )
    ).json()
    checkout_todos = [t for t in todos["items"] if t["todo_type"] == "checkout_offboarding"]
    assert checkout_todos and all(t["status"] == "completed" for t in checkout_todos)

    alerts = (
        await client.get(
            "/api/v1/workflow/alerts", headers=admin, params={"employee_id": emp["id"]}
        )
    ).json()
    assert not any(a["status"] == "active" for a in alerts["items"])

    # 台账不再包含该员工
    ledger = (await client.get("/api/v1/reports/accommodations", headers=admin)).json()
    assert emp["id"] not in {i["employee_id"] for i in ledger["items"]}


@pytest.mark.asyncio
async def test_p0_unauthenticated_and_low_priv_cannot_call_admin_apis(client: AsyncClient):
    """P0-5: 401 without token; low-privilege role gets 403 on admin APIs."""
    admin = await admin_headers(client)
    assert admin

    # unauthenticated
    r = await client.get("/api/v1/employees")
    assert r.status_code == 401
    r = await client.post("/api/v1/users", json={"username": "x", "password": "y", "display_name": "x"})
    assert r.status_code == 401

    # finance_viewer: can view reports/export but cannot manage users or allocate beds
    viewer = await make_user_with_role(client, admin, "finance_viewer", "m6fin")
    vh = await headers_for_user(client, viewer["username"])

    r = await client.post(
        "/api/v1/users",
        headers=vh,
        json={"username": "hack", "password": "HackPass12345", "display_name": "hack"},
    )
    assert r.status_code == 403, r.text

    r = await client.post(
        "/api/v1/checkin-applications",
        headers=vh,
        json={"employee_id": 1, "application_type": "new"},
    )
    assert r.status_code == 403, r.text

    r = await client.post("/api/v1/dorm-areas", headers=vh, json={"area_code": "X", "area_name": "X"})
    assert r.status_code == 403, r.text

    # allowed: report view
    r = await client.get("/api/v1/reports/accommodations", headers=vh)
    assert r.status_code == 200


@pytest.mark.asyncio
async def test_p0_excel_export_formula_injection_neutralized(client: AsyncClient):
    """P0-5: Excel export must neutralize formula injection prefixes."""
    admin = await admin_headers(client)
    assert admin
    from tests.helpers_m3 import make_employee

    emp = await make_employee(client, admin, suffix="注入")
    # attempt to store formula-like notes
    r = await client.patch(
        f"/api/v1/employees/{emp['id']}",
        headers=admin,
        json={"notes": "=cmd|'/c calc'!A1"},
    )
    assert r.status_code == 200, r.text

    export = await client.get("/api/v1/employees/export", headers=admin)
    assert export.status_code == 200
    assert export.content[:2] == b"PK"

    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(export.content))
    ws = wb.active
    formulas = []
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if isinstance(cell, str) and cell.startswith("="):
                formulas.append(cell)
    assert not formulas, formulas
